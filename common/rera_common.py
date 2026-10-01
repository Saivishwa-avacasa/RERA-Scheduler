"""
Shared plumbing for the per-state RERA scrapers.

Everything here is portal-agnostic: polite HTTP with backoff, resumable run
state, change detection against the previous run, and the Excel writer. The
state scripts (tamilnadu/, karnataka/, maharashtra/) own everything that knows what a
particular portal's HTML looks like.

Output layout, per state, mirrors maharashtra/:

    <state>/data/
        <state>_rera.xlsx   the deliverable
        snapshot.json       last completed run's records (change detection)
        changes.ndjson      append-only field-level change history
        run_log.json        last 90 run summaries
        partial/            per-source results of an unfinished run (resume state)
"""

import hashlib
import json
import re
import shutil
import time
from datetime import date, datetime, timezone
from pathlib import Path
from threading import Lock

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)

# Fields every state emits, in this order, so the workbooks line up and a later
# database sync has one shape to deal with. States may add extra fields.
COMMON_FIELDS = [
    "state_code", "rera_number", "project_name", "promoter_name",
    "promoter_normalized", "project_type", "district", "taluka", "pincode",
    "registration_date", "proposed_completion_date", "project_status",
    "detail_url", "source_url",
]

# Fields whose change is worth reporting. Bookkeeping (scraped_at, source_url,
# first_seen_at...) is excluded so an untouched project never shows as changed.
TRACKED_FIELDS = [
    "project_name", "promoter_name", "project_type", "district", "taluka",
    "pincode", "registration_date", "proposed_completion_date",
    "project_status", "status_remarks",
]

# Exit code for "finished, but some sources failed" (the rest was saved).
PARTIAL_EXIT = 4

_print_lock = Lock()


def log(msg):
    with _print_lock:
        print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def now_utc():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ─── HTTP ───────────────────────────────────────────────────────────────────

class PoliteSession:
    """requests.Session with a minimum gap between requests and retry/backoff.

    Government portals are slow and occasionally drop connections; one retry
    loop here keeps that out of every state script."""

    def __init__(self, delay=2.0, timeout=180, attempts=4):
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": USER_AGENT,
            "Accept-Language": "en-IN,en;q=0.9",
        })
        self.delay = delay
        self.timeout = timeout
        self.attempts = attempts
        self._lock = Lock()
        self._last = 0.0

    def _wait_turn(self):
        with self._lock:
            gap = self.delay - (time.monotonic() - self._last)
            if gap > 0:
                time.sleep(gap)
            self._last = time.monotonic()

    def request(self, method, url, **kwargs):
        """Returns the Response. Raises after the final attempt so callers
        record a failure rather than mistaking it for an empty result."""
        kwargs.setdefault("timeout", self.timeout)
        backoff = 3
        last_error = None
        for attempt in range(1, self.attempts + 1):
            self._wait_turn()
            try:
                response = self.session.request(method, url, **kwargs)
                if response.status_code == 200:
                    return response
                last_error = RuntimeError(f"HTTP {response.status_code} for {url}")
                # 429/5xx are worth waiting out; other 4xx will not fix themselves.
                if response.status_code not in (429, 500, 502, 503, 504):
                    raise last_error
            except requests.RequestException as exc:
                last_error = exc
            if attempt < self.attempts:
                time.sleep(backoff)
                backoff *= 2
        raise last_error

    def get(self, url, **kwargs):
        return self.request("GET", url, **kwargs)

    def post(self, url, **kwargs):
        return self.request("POST", url, **kwargs)


# ─── Normalisation ──────────────────────────────────────────────────────────

def clean(text):
    """Collapse whitespace; None-safe."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()


_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun",
     "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}

_DATE_PATTERNS = [
    # 02-01-2026, 03/01/2019, 21.12.2030 (day first, as every Indian portal writes it)
    (re.compile(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\b"), "dmy"),
    # 2026-01-02
    (re.compile(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b"), "ymd"),
    # 01-Sep-2020, 1 September 2020
    (re.compile(r"\b(\d{1,2})[-\s/]([A-Za-z]{3,9})[-\s/,]+(\d{4})\b"), "dMy"),
]


def parse_date(text):
    """First date found in `text` as ISO 'YYYY-MM-DD', or '' if none/invalid.

    Returns '' rather than raising: a bad date is a data-quality issue to
    report, not a reason to lose the rest of the record."""
    text = text or ""
    found = []
    for pattern, order in _DATE_PATTERNS:
        for match in pattern.finditer(text):
            try:
                if order == "dmy":
                    d, m, y = (int(g) for g in match.groups())
                elif order == "ymd":
                    y, m, d = (int(g) for g in match.groups())
                else:
                    d, mon, y = match.groups()
                    m = _MONTHS.get(mon[:3].lower())
                    if not m:
                        continue
                    d, y = int(d), int(y)
                found.append((match.start(), date(y, m, d).isoformat()))
            except ValueError:
                continue
    if not found:
        return ""
    return min(found)[1]


_ENTITY_SUFFIXES = [
    (r"\bPRIVATE\b", "PVT"), (r"\bPVT\.?\b", "PVT"),
    (r"\bLIMITED\b", "LTD"), (r"\bLTD\.?\b", "LTD"),
    (r"\bCOMPANY\b", "CO"), (r"\bCO\.\b", "CO"),
    (r"\bL\.?L\.?P\.?\b", "LLP"),
]


def normalize_promoter(name):
    """Canonical form used to group the same developer across spellings:
    'M/s. Newry Properties Pvt. Ltd.,' -> 'NEWRY PROPERTIES PVT LTD'."""
    name = clean(name).upper()
    name = re.sub(r"^(M/S\.?|M/S|MESSRS\.?)\s*", "", name)
    for pattern, repl in _ENTITY_SUFFIXES:
        name = re.sub(pattern, repl, name)
    name = re.sub(r"[^\w&\s]", " ", name)
    return re.sub(r"\s+", " ", name).strip()


PINCODE = re.compile(r"(?<!\d)([1-9]\d{2})\s?(\d{3})(?!\d)")


def find_pincode(text, first_digit=None):
    """Last 6-digit PIN in the text (addresses end with it). `first_digit`
    restricts to a state's PIN range, e.g. '6' for Tamil Nadu."""
    hits = ["".join(m.groups()) for m in PINCODE.finditer(text or "")]
    if first_digit:
        hits = [h for h in hits if h.startswith(first_digit)]
    return hits[-1] if hits else ""


def content_hash(record):
    payload = json.dumps({f: record.get(f, "") for f in TRACKED_FIELDS},
                         sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def validate(record, valid_states):
    """List of problems; empty means the record is usable."""
    problems = []
    if not record.get("rera_number") and not record.get("application_no"):
        problems.append("no RERA / application number")
    if not record.get("source_url"):
        problems.append("no source URL")
    if record.get("state_code") not in valid_states:
        problems.append(f"unexpected state {record.get('state_code')!r}")
    return problems


def record_key(record):
    return record.get("rera_number") or record.get("application_no")


# ─── Run state ──────────────────────────────────────────────────────────────

class RunStore:
    """Files for one state. A run is a set of named *sources* (a listing
    page, a year, a district). Each source's records are written to
    partial/<source>.json as soon as it succeeds, so an interrupted run
    resumes by skipping sources already on disk. Once a run finishes with no
    failed sources, partial/ is cleared and the next run starts fresh - which
    is how new and changed projects get picked up."""

    def __init__(self, data_dir, state_slug):
        self.dir = Path(data_dir)
        self.slug = state_slug
        self.partial = self.dir / "partial"
        self.snapshot_file = self.dir / "snapshot.json"
        self.changes_file = self.dir / "changes.ndjson"
        self.runlog_file = self.dir / "run_log.json"
        self.excel_file = self.dir / f"{state_slug}_rera.xlsx"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.partial.mkdir(exist_ok=True)

    @staticmethod
    def _safe(name):
        return re.sub(r"[^A-Za-z0-9_.-]+", "_", name)

    def done_sources(self):
        return {p.stem for p in self.partial.glob("*.json")}

    def save_source(self, name, records):
        target = self.partial / f"{self._safe(name)}.json"
        tmp = target.with_suffix(".tmp")
        tmp.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        tmp.replace(target)

    def load_partial_records(self):
        out = []
        for path in sorted(self.partial.glob("*.json")):
            try:
                out.extend(json.loads(path.read_text(encoding="utf-8")))
            except Exception:  # noqa: BLE001 - a torn file just gets re-fetched
                path.unlink(missing_ok=True)
                log(f"  discarded unreadable partial file {path.name}")
        return out

    def is_done(self, name):
        return self._safe(name) in self.done_sources()

    def clear_partial(self):
        shutil.rmtree(self.partial, ignore_errors=True)
        self.partial.mkdir(exist_ok=True)

    def load_snapshot(self):
        if self.snapshot_file.exists():
            try:
                return json.loads(self.snapshot_file.read_text(encoding="utf-8"))
            except Exception:  # noqa: BLE001
                log("Snapshot unreadable - treating this as a first run.")
        return {}

    def save_snapshot(self, records):
        self.snapshot_file.write_text(json.dumps(records, ensure_ascii=False),
                                      encoding="utf-8")

    def append_changes(self, changes):
        if not changes:
            return
        with self.changes_file.open("a", encoding="utf-8") as handle:
            for change in changes:
                handle.write(json.dumps(change, ensure_ascii=False) + "\n")

    def append_runlog(self, summary):
        history = []
        if self.runlog_file.exists():
            try:
                history = json.loads(self.runlog_file.read_text(encoding="utf-8"))
            except Exception:  # noqa: BLE001
                history = []
        history.append(summary)
        self.runlog_file.write_text(json.dumps(history[-90:], indent=2),
                                    encoding="utf-8")


def run_sources(store, sources, fetch_source, workers=1):
    """Fetch every source not already on disk. `sources` is a list of names,
    `fetch_source(name)` returns that source's records. A source that raises
    or returns nothing is recorded as failed and retried next run - it never
    stops the others."""
    from concurrent.futures import ThreadPoolExecutor, as_completed

    todo = [s for s in sources if not store.is_done(s)]
    skipped = len(sources) - len(todo)
    if skipped:
        log(f"Resuming: {skipped} of {len(sources)} sources already fetched.")
    failures = {}
    if not todo:
        return failures

    started = time.time()
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(fetch_source, name): name for name in todo}
        for index, future in enumerate(as_completed(futures), start=1):
            name = futures[future]
            try:
                records = future.result()
            except Exception as exc:  # noqa: BLE001
                failures[name] = f"{type(exc).__name__}: {str(exc)[:200]}"
                log(f"  [{index}/{len(todo)}] {name}: FAILED - {failures[name]}")
                continue
            if not records:
                # A 200 with no rows is usually a portal hiccup; recording it
                # as done would silently lose that source's projects.
                failures[name] = "returned 0 records"
                log(f"  [{index}/{len(todo)}] {name}: 0 records - will retry next run")
                continue
            store.save_source(name, records)
            log(f"  [{index}/{len(todo)}] {name}: {len(records):,} records")
    log(f"Fetched {len(todo) - len(failures)} sources in {time.time() - started:.0f}s.")
    return failures


# ─── Merge, validate, diff ──────────────────────────────────────────────────

def consolidate(raw_records, valid_states):
    """Dedupe by RERA number and split out invalid rows.

    The same registration can legitimately appear in two listings (a revised
    registration, an online and an offline list). The later-registered copy
    wins; every collision is kept in `duplicates` so it can be inspected."""
    records, invalid, duplicates = {}, [], []
    for record in raw_records:
        problems = validate(record, valid_states)
        if problems:
            invalid.append({**record, "problems": "; ".join(problems)})
            continue
        key = record_key(record)
        existing = records.get(key)
        if existing is not None:
            keep_new = ((record.get("registration_date") or "")
                        >= (existing.get("registration_date") or ""))
            kept, other = (record, existing) if keep_new else (existing, record)
            duplicates.append({"key": key, "kept_source": kept["source_url"],
                               "other_source": other["source_url"]})
            if not keep_new:
                continue
        records[key] = record
    return records, invalid, duplicates


def diff_and_stamp(previous, current, run_time, complete):
    """Compare this run against the last snapshot, carry first_seen_at and
    last_changed_at forward, and return (added, changed, removed, history).

    `removed` is only reported for a complete run: when a source failed, a
    missing project most likely just lives in that source."""
    added, changed, history = [], [], []
    for key, record in current.items():
        record["content_hash"] = content_hash(record)
        old = previous.get(key)
        if old is None:
            added.append(key)
            record["first_seen_at"] = run_time
            record["last_changed_at"] = run_time
            continue
        record["first_seen_at"] = old.get("first_seen_at", run_time)
        record["last_changed_at"] = old.get("last_changed_at", run_time)
        if old.get("content_hash") == record["content_hash"]:
            continue
        deltas = {f: (old.get(f, ""), record.get(f, ""))
                  for f in TRACKED_FIELDS if old.get(f, "") != record.get(f, "")}
        if not deltas:
            continue
        record["last_changed_at"] = run_time
        changed.append({"key": key, "changes": deltas})
        for field, (before, after) in deltas.items():
            history.append({"changed_at": run_time, "key": key, "field": field,
                            "old_value": before, "new_value": after,
                            "source_url": record.get("source_url", "")})

    removed = [k for k in previous if k not in current] if complete else []
    if not complete:
        # Keep last-known copies of projects we could not re-check this time.
        for key, old in previous.items():
            current.setdefault(key, old)
    return added, changed, removed, history


# ─── Excel ──────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FCE4E4")


def _sheet(workbook, title, columns, rows, flag=None):
    sheet = workbook.create_sheet(title)
    for index, (_, label, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append([_cell(row.get(key, "")) for key, _, _ in columns])
        if flag and flag(row):
            for column in range(1, len(columns) + 1):
                sheet.cell(row=sheet.max_row, column=column).fill = FLAG_FILL
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
    return sheet


def _cell(value):
    if isinstance(value, (list, dict, tuple)):
        value = json.dumps(value, ensure_ascii=False)
    value = "" if value is None else value
    # Excel rejects control characters and caps cells at 32,767 chars.
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32000]
    return value


def write_workbook(path, project_columns, records, *, flag=None, changes=(),
                   invalid=(), duplicates=(), extra_sheets=(), summary=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    ordered = sorted(records.values(), key=lambda r: record_key(r) or "")
    _sheet(workbook, "Projects", project_columns, ordered, flag)
    if flag:
        _sheet(workbook, "Flagged", project_columns,
               [r for r in ordered if flag(r)])
    _sheet(workbook, "Changes This Run",
           [("changed_at", "Changed At", 22), ("key", "RERA / App No", 30),
            ("field", "Field", 24), ("old_value", "Old", 40),
            ("new_value", "New", 40), ("source_url", "Source", 50)],
           list(changes))
    _sheet(workbook, "Invalid Rows",
           [("problems", "Problems", 40)] + project_columns, list(invalid))
    _sheet(workbook, "Duplicates",
           [("key", "RERA / App No", 30), ("kept_source", "Kept (source)", 60),
            ("other_source", "Also listed in", 60)], list(duplicates))
    for title, columns, rows in extra_sheets:
        _sheet(workbook, title, columns, rows)
    if summary:
        _sheet(workbook, "Run Summary", [("field", "Field", 32), ("value", "Value", 80)],
               [{"field": k, "value": str(v)} for k, v in summary.items()])
    workbook.save(path)


# ─── Finish a run ───────────────────────────────────────────────────────────

def finish_run(store, *, state_code, raw_records, failures, started,
               project_columns, flag=None, extra_sheets=(), extra_summary=None,
               valid_states=None, sources_total=None, partial_scope=False):
    """Everything after fetching: consolidate, diff, write Excel/snapshot/
    history/run log, print the outcome. Returns a process exit code.

    `partial_scope` marks a deliberately limited run (--limit, --district):
    projects outside that scope are carried over from the last snapshot
    instead of being reported as no longer listed."""
    valid_states = valid_states or {state_code}
    run_time = now_utc()
    records, invalid, duplicates = consolidate(raw_records, valid_states)
    if not records:
        log("No valid project rows collected - nothing written.")
        return 1

    complete = not failures
    scraped_now = dict(records)  # before diff_and_stamp carries old rows forward
    previous = store.load_snapshot()
    added, changed, removed, history = diff_and_stamp(
        previous, records, run_time, complete and not partial_scope)
    flagged = sum(1 for r in records.values() if flag and flag(r))
    summary = {
        "state": state_code,
        "run_started_utc": started.isoformat(timespec="seconds"),
        "duration_seconds": round((datetime.now(timezone.utc) - started).total_seconds(), 1),
        "sources_total": sources_total if sources_total is not None else "",
        "sources_failed": len(failures),
        "rows_scraped": len(raw_records),
        "total_projects": len(records),
        "invalid_rows": len(invalid),
        "duplicate_rows": len(duplicates),
        "flagged_projects": flagged,
        "new_since_last_run": len(added),
        "changed_since_last_run": len(changed),
        "disappeared_since_last_run": (len(removed) if complete and not partial_scope
                                       else "not computed (partial or incomplete run)"),
        "complete": complete,
        **(extra_summary or {}),
    }
    if failures:
        summary["failed_sources"] = "; ".join(f"{k} ({v})" for k, v in failures.items())

    log("Writing Excel...")
    write_workbook(store.excel_file, project_columns, records, flag=flag,
                   changes=history, invalid=invalid, duplicates=duplicates,
                   extra_sheets=extra_sheets, summary=summary)
    store.save_snapshot(records)
    store.append_changes(history)
    store.append_runlog(summary)
    if complete:
        store.clear_partial()

    db_error = sync_to_database(state_code, scraped_now, started, summary, failures)

    print()
    if not previous:
        log(f"No local snapshot (first run, or a fresh container): {len(records):,} "
            f"projects scraped. The database line reports what was new or changed.")
    elif not (added or changed or removed):
        log(f"Already up to date - {len(records):,} projects, nothing changed.")
    else:
        log(f"Updated: {len(added):,} new, {len(changed):,} changed, "
            f"{len(removed):,} no longer listed.")
        for entry in changed[:5]:
            log(f"    {entry['key']}: {', '.join(entry['changes'])}")
    if invalid:
        log(f"{len(invalid):,} rows failed validation - see the 'Invalid Rows' sheet.")
    if duplicates:
        log(f"{len(duplicates):,} duplicate listings merged - see the 'Duplicates' sheet.")
    if flag:
        log(f"{flagged:,} projects carry a negative status.")
    if failures:
        log(f"WARNING: {len(failures)} sources failed; re-run the same command to "
            f"retry only those: {', '.join(failures)}")
    log(f"Excel: {store.excel_file}")
    if db_error:
        return 3
    # 4, not 2: Python itself exits 2 for "can't open file" and argparse errors,
    # and those must never look like a partly successful run.
    return 0 if complete else PARTIAL_EXIT


def sync_to_database(state_code, records, started, summary, failures):
    """Push this run to DATABASE_URL if one is configured. Returns an error
    message (and logs it) instead of raising, so a database outage never
    loses the Excel output that was already written."""
    from common import db

    if not db.database_url():
        log("DATABASE_URL not set - skipping database sync (Excel/JSON only).")
        return None
    try:
        db.sync_run(state_code, records, started=started, summary=summary, failures=failures)
        return None
    except Exception as exc:  # noqa: BLE001
        # Never echo the URL: it carries the password.
        message = db.redact(f"{type(exc).__name__}: {str(exc).splitlines()[0][:300]}")
        log(f"ERROR: database sync failed - {message}")
        return message
